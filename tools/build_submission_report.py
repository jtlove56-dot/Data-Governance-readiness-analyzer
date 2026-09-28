from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOWNLOADS = Path("/Users/charlesnwabuike/Downloads")
ASSET_DIR = Path("/tmp/dgra_report_assets")
OUTPUT = DOWNLOADS / "Milestone2_Group_Report_APA_Submission_Ready.docx"

NAVY = "1F3A4D"
PALE = "EAF0F3"
MID = "65727A"
BLACK = RGBColor(0, 0, 0)


def crop_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    sources = {
        "describe": (ASSET_DIR / "01_describe.png", (0, 0, 633, 510)),
        "questions": (ASSET_DIR / "02_questions_completed.png", (0, 0, 633, 854)),
        "assessment": (ASSET_DIR / "03_assessment.png", (0, 0, 633, 550)),
        "recommendations": (ASSET_DIR / "04_recommendations.png", (0, 0, 633, 735)),
    }
    output: dict[str, Path] = {}
    for name, (source, box) in sources.items():
        target = ASSET_DIR / f"{name}_cropped.png"
        with Image.open(source) as image:
            image.crop(box).convert("RGB").save(target, "PNG", optimize=True)
        output[name] = target

    jira_source = DOWNLOADS / "DGRA_Jira_Backlog_Clean.png"
    jira_overview = ASSET_DIR / "jira_backlog_overview.png"
    jira_detail = ASSET_DIR / "jira_sprints_detail.png"
    with Image.open(jira_source) as image:
        image.convert("RGB").save(jira_overview, "PNG", optimize=True)
        image.crop((1220, 300, 2990, 1640)).convert("RGB").save(
            jira_detail, "PNG", optimize=True
        )
    output["jira_overview"] = jira_overview
    output["jira_detail"] = jira_detail
    return output


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, indent: bool = True) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.first_line_indent = Inches(0.5) if indent else None


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_figure(
    doc: Document,
    number: int,
    title: str,
    image_path: Path,
    note: str,
    width: float = 6.45,
) -> None:
    number_paragraph = doc.add_paragraph()
    number_paragraph.paragraph_format.keep_with_next = True
    number_paragraph.paragraph_format.line_spacing = 1
    number_paragraph.add_run(f"Figure {number}").bold = True
    title_paragraph = doc.add_paragraph()
    title_paragraph.paragraph_format.keep_with_next = True
    title_paragraph.paragraph_format.line_spacing = 1
    title_run = title_paragraph.add_run(title)
    title_run.italic = True
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.line_spacing = 1
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    note_paragraph = doc.add_paragraph()
    note_paragraph.paragraph_format.line_spacing = 1
    note_paragraph.paragraph_format.space_after = Pt(12)
    note_paragraph.add_run("Note. ").italic = True
    note_paragraph.add_run(note)


def add_status_table(doc: Document) -> None:
    data = [
        ("Sprint 1: Foundation", "SCRUM-6–SCRUM-9", "23", "4 Done"),
        ("Sprint 2: Assessment", "SCRUM-10–SCRUM-14", "28", "2 Done; 3 To Do"),
        ("Sprint 3: Demo Ready", "SCRUM-15–SCRUM-18", "21", "2 Done; 1 In Progress; 1 To Do"),
        ("Total", "13 issues", "72", "8 Done; 1 In Progress; 4 To Do"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headings = ["Sprint", "Issue range", "Points", "Status at milestone snapshot"]
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, label in enumerate(headings):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].paragraph_format.line_spacing = 1
    for row_index, row_data in enumerate(data):
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, PALE)
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(10)
            cell.paragraphs[0].paragraph_format.line_spacing = 1
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(1.4)
    table.columns[2].width = Inches(0.7)
    table.columns[3].width = Inches(2.3)
    note = doc.add_paragraph()
    note.paragraph_format.line_spacing = 1
    note.add_run("Note. ").italic = True
    note.add_run(
        "Status values reproduce the September 26, 2026 milestone snapshot in the supplied report; "
        "Jira remains the source of truth for later changes."
    )


def add_verification_table(doc: Document) -> None:
    rows = [
        ("Backend test suite", "31 passed", "Pass"),
        ("Frontend unit tests", "25 passed", "Pass"),
        ("Static analysis", "ESLint and TypeScript checks completed", "Pass"),
        ("Production build", "Next.js optimized build completed", "Pass"),
        ("Guided workflow", "Describe → questionnaire → score → recommendations", "Pass"),
        ("Representative output", "Medium risk, 50/100, with four contributing factors", "Observed"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, label in enumerate(("Verification", "Evidence", "Result")):
        set_cell_shading(header.cells[idx], NAVY)
        set_cell_margins(header.cells[idx])
        run = header.cells[idx].paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        header.cells[idx].paragraphs[0].paragraph_format.line_spacing = 1
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(values):
            set_cell_margins(row.cells[idx])
            if row_index % 2 == 0:
                set_cell_shading(row.cells[idx], PALE)
            run = row.cells[idx].paragraphs[0].add_run(value)
            run.font.size = Pt(10)
            row.cells[idx].paragraphs[0].paragraph_format.line_spacing = 1
            if idx == 2:
                run.bold = True


def add_story(doc: Document, story: dict) -> None:
    add_heading(doc, story["heading"], 3)
    metadata = doc.add_paragraph()
    metadata.paragraph_format.keep_with_next = True
    metadata_run = metadata.add_run(story["meta"])
    metadata_run.bold = True
    metadata_run.font.color.rgb = RGBColor.from_string(MID)
    add_body(doc, story["story"])
    label = doc.add_paragraph()
    label.paragraph_format.keep_with_next = True
    label.add_run("Acceptance criteria").bold = True
    add_bullets(doc, story["criteria"])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    add_page_number(section.header.paragraphs[0])

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_border = title._element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        title._element.pPr.remove(title_border)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    styles["List Bullet"].font.name = "Times New Roman"
    styles["List Bullet"].font.size = Pt(12)
    styles["List Bullet"].paragraph_format.line_spacing = 2
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.5)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.25)

    if "Figure Caption" not in styles:
        styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)


def build_report() -> Path:
    images = crop_assets()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Milestone 2 Group Report: Data Governance Readiness Analyzer"
    doc.core_properties.author = (
        "Jelani Denmark; Chidiebube Nwabuike; Ademide Ogunmefun; Jt Love"
    )
    doc.core_properties.subject = "Senior Capstone Milestone 2 Group Report"
    doc.core_properties.keywords = "data governance, privacy risk, Jira, capstone"

    # APA-style student title page.
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Milestone 2 Group Report: Data Governance Readiness Analyzer")
    for text in (
        "Jelani Denmark, Chidiebube Nwabuike, Ademide Ogunmefun, and Jt Love",
        "Computer Science, Bowie State University",
        "COSC 480-001: Senior Capstone",
        "Ruth Olusegun",
        "September 27, 2026",
    ):
        paragraph = doc.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    main_title = doc.add_paragraph(style="Title")
    main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    main_title.add_run("Milestone 2 Group Report: Data Governance Readiness Analyzer")
    add_heading(doc, "Executive Summary", 1)
    add_body(
        doc,
        "The Data Governance Readiness Analyzer is a guided web application that helps an "
        "organization evaluate whether a proposed data-sharing initiative is ready to proceed. "
        "A user describes the initiative in ordinary language, answers a structured governance "
        "questionnaire, and receives an explainable privacy risk score, the factors behind that "
        "score, and safeguards mapped to privacy-preserving capabilities for zero-trust data exchange "
        "(Data Governance Readiness Analyzer Team, 2026a).",
    )
    add_body(
        doc,
        "The implementation uses a Next.js frontend and a FastAPI backend. The workflow is "
        "organized into four visible steps: Describe, Answer Questions, Assessment, and "
        "Recommendations. It applies a deterministic 0–100 scoring rubric. This submission "
        "combines the milestone backlog, testable user stories, current implementation evidence, "
        "and an independently rerun verification summary."
    )

    add_heading(doc, "Project Scope and Backlog", 1)
    add_body(
        doc,
        "All milestone work is tracked under parent epic SCRUM-5 and organized across three "
        "delivery increments: Sprint 1 (Foundation), Sprint 2 (Assessment), and Sprint 3 (Demo "
        "Ready) (Data Governance Readiness Analyzer Team, 2026b). The supplied milestone "
        "snapshot contains 13 issues (SCRUM-6 through SCRUM-18) totaling 72 story points."
    )
    doc.add_page_break()
    table_number = doc.add_paragraph()
    table_number.paragraph_format.keep_with_next = True
    table_number.paragraph_format.line_spacing = 1
    table_number.add_run("Table 1").bold = True
    table_title = doc.add_paragraph()
    table_title.paragraph_format.keep_with_next = True
    table_title.paragraph_format.line_spacing = 1
    table_title.add_run("Milestone backlog summary").italic = True
    add_status_table(doc)
    add_figure(
        doc,
        1,
        "Data Governance Readiness Analyzer Jira backlog",
        images["jira_overview"],
        "Live, read-only backlog view captured September 27, 2026. The screen shows the epic "
        "panel, completed Sprint 1 work, and Sprint 2 status. Source: Data Governance Readiness "
        "Analyzer Team (2026b).",
    )

    doc.add_page_break()
    add_heading(doc, "Epics", 1)
    add_heading(doc, "Epic 1: Foundation and Guided Intake Flow", 2)
    add_body(
        doc,
        "Sprint 1 establishes the product rules and technical foundation. It confirms a "
        "versioned scoring rubric and approved capability language, establishes the Next.js and "
        "FastAPI application architecture, and builds the guided intake and structured "
        "questionnaire. Its four issues, SCRUM-6 through SCRUM-9, total 23 points and are marked Done."
    )
    add_heading(doc, "Epic 2: Privacy Risk Assessment and Safeguard Mapping", 2)
    add_body(
        doc,
        "Sprint 2 turns questionnaire answers into an explainable decision-support result. A "
        "deterministic rules service returns a score from 0 to 100, classifies it as Low (0–29), "
        "Medium (30–59), or High (60–100), explains each contributor, and maps risks to "
        "safeguards. SCRUM-10 through SCRUM-14 total 28 points."
    )
    add_heading(doc, "Epic 3: End-to-End Experience, Export, and Deployment", 2)
    add_body(
        doc,
        "Sprint 3 prepares the MVP for demonstration and release. It covers a stakeholder-ready "
        "PDF report, privacy-conscious session handling, accessibility and usability checks, and "
        "end-to-end deployment verification. SCRUM-15 through SCRUM-18 total 21 points."
    )
    add_figure(
        doc,
        2,
        "Sprint 1 and Sprint 2 issue status detail",
        images["jira_detail"],
        "Enlarged detail from the live Jira backlog capture. The milestone statuses reported in "
        "this document are tied to the September 26 snapshot in the source draft."
    )

    doc.add_page_break()
    add_heading(doc, "User Stories With Acceptance Criteria", 1)
    add_body(
        doc,
        "The ten user stories below retain the issue ownership, point values, milestone status, "
        "and testable criteria from the supplied report. SCRUM-7 (architecture), SCRUM-12 "
        "(recommendation mapping), and SCRUM-18 (end-to-end testing and deployment) are "
        "supporting technical tasks summarized in their related epics.",
        indent=False,
    )

    epic_1 = [
        {
            "heading": "US-1 (SCRUM-6): Confirmed Scoring Rubric and Capability Language",
            "meta": "Owner: Chidiebube Nwabuike | 5 points | Done",
            "story": "As a product stakeholder, I want a versioned risk-scoring rubric and approved capability language so that every score and recommendation is consistent, explainable, and accurate.",
            "criteria": [
                "Every questionnaire response maps to a documented scoring rule.",
                "Low, Medium, and High thresholds are explicit and reviewable.",
                "Each recommended capability has approved plain-language copy and applicability criteria.",
            ],
        },
        {
            "heading": "US-2 (SCRUM-8): Guided Assessment and Use-Case Description",
            "meta": "Owner: Chidiebube Nwabuike | 5 points | Done",
            "story": "As a non-specialist evaluating a data initiative, I want to describe my proposed use case in plain language and see a clear assessment process so that I can begin without understanding privacy terminology.",
            "criteria": [
                "Empty or invalid submissions receive clear corrective guidance.",
                "The current step and remaining steps are always visible.",
                "Navigating backward preserves previously entered data.",
            ],
        },
        {
            "heading": "US-3 (SCRUM-9): Structured Governance Questionnaire",
            "meta": "Owner: Ademide Ogunmefun | 8 points | Done",
            "story": "As a non-specialist, I want short, understandable follow-up questions about my data use case so that the system can assess risk without requiring privacy expertise.",
            "criteria": [
                "Every scoring input required by the approved rubric is collected.",
                "Required questions cannot be skipped without clear validation.",
                "Submitted answers match the backend assessment request schema.",
            ],
        },
    ]
    add_heading(doc, "Epic 1: Foundation and Guided Intake Flow", 2)
    for story in epic_1:
        add_story(doc, story)
    add_figure(
        doc,
        3,
        "Guided use-case description step",
        images["describe"],
        "The running application presents plain-language guidance, visible progress, and a "
        "privacy notice before collecting the proposed use case."
    )
    add_figure(
        doc,
        4,
        "Completed structured governance questionnaire",
        images["questions"],
        "The representative run selected Names and Email addresses, third-party access, no raw "
        "identifier exchange, no uncontrolled transfer, dataset combination, and no reuse beyond "
        "the stated purpose.",
        width=4.5,
    )

    doc.add_page_break()
    add_heading(doc, "Epic 2: Privacy Risk Assessment and Safeguard Mapping", 2)
    epic_2 = [
        {
            "heading": "US-4 (SCRUM-10): Rule-Based Risk-Scoring Service",
            "meta": "Owner: Jt Love | 8 points | Done",
            "story": "As a user completing an assessment, I want my answers scored by consistent, rule-based logic so that I receive the same trustworthy result every time for the same inputs.",
            "criteria": [
                "Valid responses produce a score from 0 through 100 with a level matching the approved thresholds.",
                "Identical inputs produce identical results for the same rules version.",
                "Invalid input returns a clear client error and never a misleading score.",
            ],
        },
        {
            "heading": "US-5 (SCRUM-11): Risk Score, Level, and Contributing Factors",
            "meta": "Owner: Jt Love | 5 points | Done",
            "story": "As a user completing an assessment, I want to understand both the risk level and why the system assigned it so that I can make an informed decision.",
            "criteria": [
                "The displayed score and level match the assessment API response.",
                "Each contributing factor has a clear plain-language explanation.",
                "The screen states that the assessment is decision support, not legal advice.",
            ],
        },
        {
            "heading": "US-6 (SCRUM-13 with SCRUM-12): Tailored Recommendations",
            "meta": "Owner: Unassigned | 10 combined points | To Do at milestone snapshot",
            "story": "As a user reviewing an assessment, I want prioritized, understandable safeguards, including relevant privacy-preserving guidance, so that I know what action to take next.",
            "criteria": [
                "Each recommendation includes a concise rationale traceable to an identified risk factor.",
                "Capability guidance appears only when documented applicability conditions are met and uses approved language.",
                "General safeguards remain available when a specialized capability is not the appropriate fit.",
            ],
        },
        {
            "heading": "US-7 (SCRUM-14): Verified Correctness and Response Time",
            "meta": "Owner: Unassigned | 5 points | To Do at milestone snapshot",
            "story": "As a project stakeholder, I want automated evidence that scoring is correct and fast so that we can trust the results shown in the demo.",
            "criteria": [
                "Automated tests verify all approved scoring thresholds and exercise every supported risk factor.",
                "The assessment API responds within 2 seconds under documented test conditions.",
                "Test results are reproducible from documented commands.",
            ],
        },
    ]
    for story in epic_2:
        add_story(doc, story)
    add_figure(
        doc,
        5,
        "Explainable privacy risk result",
        images["assessment"],
        "The representative assessment returned 50/100 (Medium). The display identifies four "
        "contributors totaling 50 points and communicates the approval condition in plain language."
    )

    doc.add_page_break()
    add_heading(doc, "Epic 3: End-to-End Experience, Export, and Deployment", 2)
    epic_3 = [
        {
            "heading": "US-8 (SCRUM-15): Professional Assessment Report",
            "meta": "Owner: Jelani Denmark | 8 points | Done",
            "story": "As a user or governance reviewer, I want to download a clear assessment report so that I can share the use case, findings, and recommended next steps with stakeholders.",
            "criteria": [
                "Users can generate and download a PDF from a completed assessment.",
                "The report accurately reflects the assessment shown in the application, with all required sections labeled.",
                "The report contains no internal identifiers or unnecessary sensitive metadata.",
            ],
        },
        {
            "heading": "US-9 (SCRUM-16): Privacy-Conscious Data Handling",
            "meta": "Owner: Jelani Denmark | 5 points | Done",
            "story": "As a user assessing a sensitive use case, I want the tool to collect and keep only what it needs so that my own data is handled by the same minimization principle the tool recommends.",
            "criteria": [
                "Every collected field has a documented assessment purpose.",
                "Sensitive assessment content is excluded from routine application logs.",
                "Users receive accurate information about whether their assessment is retained.",
            ],
        },
        {
            "heading": "US-10 (SCRUM-17): Usable, Accessible, Plain-Language Experience",
            "meta": "Owner: Jelani Denmark | 3 points | In Progress at milestone snapshot",
            "story": "As a non-specialist, I want the assessment to be understandable and accessible so that I can complete it in under 5 minutes without separate instructions.",
            "criteria": [
                "The median or representative completion time is under 5 minutes.",
                "All interactive controls are keyboard operable with visible focus.",
                "No unresolved high-severity usability or accessibility defects remain for the demo.",
            ],
        },
    ]
    for story in epic_3:
        add_story(doc, story)
    add_figure(
        doc,
        6,
        "Tailored safeguards and privacy-preserving capability fit",
        images["recommendations"],
        "The final step translates the identified risks into six actionable controls, states the "
        "privacy-preserving capability fit, and exposes the assessment-record and PDF-report controls.",
        width=5.2,
    )

    doc.add_page_break()
    add_heading(doc, "Implementation Verification", 1)
    add_body(
        doc,
        "The repository was rerun on September 27, 2026. The backend and frontend test suites, "
        "static checks, and production build all completed successfully. The application was then "
        "exercised end to end with a representative record-matching scenario. No source-code "
        "changes were required to produce the evidence in this report (Data Governance Readiness "
        "Analyzer Team, 2026c)."
    )
    add_verification_table(doc)
    add_body(
        doc,
        "The representative scenario involved names and email addresses, access by another "
        "organization, dataset combination, and customer or record matching. The engine returned "
        "50 points: direct identifiers (+14), access by another organization (+18), "
        "re-identification potential (+12), and the intended use (+6). This result demonstrates "
        "that the visible score is traceable to the rubric rather than presented as an unexplained rating."
    )

    add_heading(doc, "Team Workflow", 1)
    add_body(
        doc,
        "Chidiebube Nwabuike established the Jira Scrum board on September 14, 2026, creating "
        "parent epic SCRUM-5 and issues SCRUM-6 through SCRUM-18 from the product requirements "
        "document. Each issue uses a consistent structure that includes context or a user story, "
        "requirements, testable acceptance criteria, dependencies, sprint placement, and a point "
        "estimate. This structure supports traceability from planned work to implementation evidence."
    )
    add_body(
        doc,
        "Responsibility is divided by product area. Chidiebube owns the rubric, architecture, and "
        "guided-flow shell; Ademide Ogunmefun owns the questionnaire; Jt Love owns the scoring "
        "service and results view; and Jelani Denmark owns the PDF report, data handling, and "
        "accessibility validation. Jira remains the shared coordination record for status and ownership."
    )

    add_heading(doc, "Current Limitations and Next Steps", 1)
    add_bullets(
        doc,
        [
            "Complete the remaining Jira items and reconcile milestone labels with the live board before the next review.",
            "Record the documented performance conditions for the under-two-second API acceptance criterion.",
            "Complete keyboard, focus-visibility, and representative completion-time evidence for SCRUM-17.",
            "Run the final deployment check and attach the production URL and release evidence to SCRUM-18.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "References", 1)
    references = [
        (
            "Data Governance Readiness Analyzer Team. (2026a, August 31). Product requirements "
            "document: Data Governance Readiness Analyzer (Version 1.0 draft) [Unpublished internal document]."
        ),
        (
            "Data Governance Readiness Analyzer Team. (2026b). SCRUM project backlog [Jira board]. "
            "Atlassian. Retrieved September 27, 2026, from "
            "https://datagovanalyzer.atlassian.net/jira/software/projects/SCRUM/boards/1/backlog"
        ),
        (
            "Data Governance Readiness Analyzer Team. (2026c). Data Governance Readiness Analyzer "
            "source code and automated test suite [Computer software]."
        ),
    ]
    for reference in references:
        paragraph = doc.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
